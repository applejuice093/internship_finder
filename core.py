"""Pure logic for the internship recommender (no Streamlit imports).

Keeping this separate from app.py means the parsing, skill extraction
and matching code can be unit tested without starting a Streamlit app.
"""
import json
import re
from io import BytesIO

import docx
import pdfplumber
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# Skills we look for in a resume, grouped only for readability.
TARGET_SKILLS = {
    "Core_Languages": ["python", "c++", "c", "java", "javascript", "sql", "typescript"],
    "AI_ML": ["pytorch", "tensorflow", "scikit-learn", "pandas", "llm"],
    "Frameworks": ["fastapi", "react", "node", "django", "spring"],
    "Daily_Tooling": ["git", "github", "docker", "aws"],
}

# Built-in sample postings used as the internship "database".
BUILTIN_INTERNSHIPS = [
    {"id": 1, "title": "Junior Backend Intern", "company": "TechGram", "location": "Remote", "education": "Graduate", "skills": ["python", "sql", "api"], "description": "Build robust backend APIs using Python and SQL for rural e-governance platforms.", "apply_url": "https://pminternship.mca.gov.in/"},
    {"id": 2, "title": "Frontend Trainee", "company": "JanMitra", "location": "Maharashtra", "education": "Diploma", "skills": ["javascript", "react", "ui"], "description": "Design accessible, user-friendly interfaces using React and JavaScript for mobile devices.", "apply_url": "https://pminternship.mca.gov.in/"},
    {"id": 3, "title": "Data & Cloud Intern", "company": "AgriData Co.", "location": "Remote", "education": "Graduate", "skills": ["python", "aws", "sql"], "description": "Manage massive agricultural crop databases hosted on AWS cloud.", "apply_url": "https://pminternship.mca.gov.in/"},
]

# Questions used when no Gemini API key is configured (offline mode).
OFFLINE_QUESTIONS = [
    {"q": "What is Object-Oriented Programming (OOP)?", "options": ["A programming paradigm", "A hardware component", "A database structure", "An operating system"], "a": "A programming paradigm"},
    {"q": "Which data structure uses a Last-In-First-Out (LIFO) model?", "options": ["Queue", "Array", "Stack", "Linked List"], "a": "Stack"},
    {"q": "What is the primary purpose of version control systems like Git?", "options": ["Compiling code", "Tracking code changes and collaboration", "Designing user interfaces", "Scanning for viruses"], "a": "Tracking code changes and collaboration"},
]


def build_question_prompt(skills, num_questions):
    """Prompt asking the model for multiple-choice questions as JSON."""
    return f"""
        You are a technical interviewer. The candidate has these skills: {', '.join(skills)}.
        Generate {num_questions} UNIQUE and DISTINCT multiple-choice interview questions testing different aspects of these skills.
        Ensure no two questions are about the same exact topic.
        Return ONLY a valid JSON array of objects. Do not use markdown blocks.
        Format: [{{"q": "Question text?", "options": ["A", "B", "C", "D"], "a": "Exact text of correct option"}}]
        """


def _is_valid_question(item):
    """A usable question has text, 2+ unique string options and an answer among them."""
    if not isinstance(item, dict):
        return False
    question, options, answer = item.get("q"), item.get("options"), item.get("a")
    return (
        isinstance(question, str)
        and question.strip() != ""
        and isinstance(options, list)
        and len(options) >= 2
        and all(isinstance(option, str) for option in options)
        and len(set(options)) == len(options)
        and answer in options
    )


def parse_questions(response_text):
    """Turn the model's text reply into a list of valid question dicts.

    Models sometimes wrap the JSON in markdown fences or add a sentence
    before it, and occasionally return a question whose answer is not one
    of its options (which would make it impossible to answer correctly).
    Those questions are dropped. Raises ValueError if nothing usable is left.
    """
    clean = response_text.replace("```json", "").replace("```", "").strip()
    start, end = clean.find("["), clean.rfind("]")
    if start == -1 or end < start:
        raise ValueError("No JSON array found in the model response")
    try:
        items = json.loads(clean[start:end + 1])
    except json.JSONDecodeError as exc:
        raise ValueError(f"Model response is not valid JSON: {exc}") from exc
    questions = [item for item in items if _is_valid_question(item)]
    if not questions:
        raise ValueError("Model response contained no valid questions")
    return questions


SUPPORTED_RESUME_TYPES = ["pdf", "docx"]


def read_resume_text(filename, file_bytes):
    """Extract plain text from an uploaded PDF or Word (.docx) resume.

    The extension check is case-insensitive, so "Resume.PDF" works too.
    Raises ValueError for any other file type.
    """
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "pdf":
        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    if extension == "docx":
        document = docx.Document(BytesIO(file_bytes))
        lines = [paragraph.text for paragraph in document.paragraphs]
        # Many resume templates put skills in tables.
        for table in document.tables:
            for row in table.rows:
                lines.extend(cell.text for cell in row.cells)
        return "\n".join(lines)
    raise ValueError(f"Unsupported resume type: {filename!r}")


def _skill_pattern(skill):
    r"""Regex that matches skill as a whole token.

    Plain \b boundaries fail for symbols: r"\bc\+\+\b" never matches
    "C++ " (there is no word boundary after "+"), and r"\bc\b" matches the
    "c" inside "C++", "C#" and "Objective-C". Lookarounds that also treat
    "+", "#" and "-" as part of a token fix both cases.
    """
    return re.compile(r"(?<![\w+#-])" + re.escape(skill) + r"(?![\w+#])")


def extract_skills_from_text(text):
    """Return the TARGET_SKILLS mentioned in text, in TARGET_SKILLS order."""
    found = []
    text_lower = text.lower()
    for skills in TARGET_SKILLS.values():
        for skill in skills:
            if skill not in found and _skill_pattern(skill).search(text_lower):
                found.append(skill)
    return found


def rag_semantic_search(user_profile, user_skills, db_jobs, top_n=3):
    """Rank internships by TF-IDF cosine similarity to the candidate profile."""
    corpus = [f"{j['title']} {j['company']} {j['location']} {' '.join(j['skills'])} {j['description']}" for j in db_jobs]
    query = f"{user_profile['location']} {user_profile['education']} {' '.join(user_skills)}"

    vectorizer = TfidfVectorizer()
    tfidf_matrix = vectorizer.fit_transform(corpus + [query])
    cosine_similarities = cosine_similarity(tfidf_matrix[-1], tfidf_matrix[:-1]).flatten()

    ranked = []
    for idx, score in enumerate(cosine_similarities):
        intern_data = db_jobs[idx].copy()
        intern_data['score'] = min(int(score * 100) + 15, 99)
        ranked.append(intern_data)

    ranked.sort(key=lambda x: x['score'], reverse=True)
    return ranked[:top_n]
