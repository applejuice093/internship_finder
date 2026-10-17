from io import BytesIO

import docx
import pytest

from core import extract_skills_from_text, read_resume_text


def make_pdf(text):
    """Build a minimal one-page PDF containing text (no extra dependencies)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraph, table_cell=None):
    document = docx.Document()
    document.add_paragraph(paragraph)
    if table_cell:
        document.add_table(rows=1, cols=1).cell(0, 0).text = table_cell
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_reads_pdf():
    text = read_resume_text("resume.pdf", make_pdf("Skills: Python and SQL"))
    assert extract_skills_from_text(text) == ["python", "sql"]


def test_pdf_extension_is_case_insensitive():
    assert "Docker" in read_resume_text("Resume.PDF", make_pdf("Docker"))


def test_reads_docx_paragraphs_and_tables():
    data = make_docx("Experienced with Java", table_cell="React, AWS")
    text = read_resume_text("cv.docx", data)
    assert extract_skills_from_text(text) == ["java", "react", "aws"]


@pytest.mark.parametrize("name", ["resume.txt", "resume", "photo.png"])
def test_rejects_unsupported_types(name):
    with pytest.raises(ValueError):
        read_resume_text(name, b"")
